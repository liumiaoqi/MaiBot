# AI生成
"""
方案一：PPT → PDF
=================
1. 用 python-pptx 生成 .pptx 文件（A4 纵向尺寸）
2. 用 comtypes 调用 PowerPoint COM 接口转为 PDF

用法：
    .venv\\Scripts\\python.exe ppt_to_pdf.py

输出：
    output_ppt.pptx  →  output_ppt.pdf
"""

import os
import sys

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# 导入共享内容
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import CONTENT

# ── 配置 ──
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PPTX_PATH = os.path.join(BASE_DIR, "output_ppt.pptx")
PDF_PATH = os.path.join(BASE_DIR, "output_ppt.pdf")

# A4 尺寸（纵向）
A4_WIDTH = Cm(21)
A4_HEIGHT = Cm(29.7)

# 颜色
COLOR_TITLE = RGBColor(0x1A, 0x1A, 0x2E)
COLOR_BODY = RGBColor(0x33, 0x33, 0x33)
COLOR_ACCENT = RGBColor(0x4A, 0x90, 0xD9)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_HEADER_BG = RGBColor(0x4A, 0x90, 0xD9)


def add_title_slide(prs, page):
    """添加封面页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局

    # 标题
    txBox = slide.shapes.add_textbox(Cm(3), Cm(10), Cm(15), Cm(3))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = page["title"]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE
    p.alignment = PP_ALIGN.CENTER

    # 副标题
    txBox2 = slide.shapes.add_textbox(Cm(3), Cm(14), Cm(15), Cm(2))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = page.get("subtitle", "")
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_ACCENT
    p2.alignment = PP_ALIGN.CENTER


def add_text_slide(prs, page):
    """添加文本页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题
    txBox = slide.shapes.add_textbox(Cm(2), Cm(1.5), Cm(17), Cm(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = page["title"]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE

    # 正文
    if page.get("body"):
        txBox2 = slide.shapes.add_textbox(Cm(2), Cm(4.5), Cm(17), Cm(3))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = page["body"]
        p2.font.size = Pt(16)
        p2.font.color.rgb = COLOR_BODY

    # 要点列表
    if page.get("bullets"):
        txBox3 = slide.shapes.add_textbox(Cm(2.5), Cm(8.5), Cm(16), Cm(15))
        tf3 = txBox3.text_frame
        tf3.word_wrap = True
        for i, bullet in enumerate(page["bullets"]):
            if i == 0:
                p3 = tf3.paragraphs[0]
            else:
                p3 = tf3.add_paragraph()
            p3.text = f"• {bullet}"
            p3.font.size = Pt(16)
            p3.font.color.rgb = COLOR_BODY
            p3.space_after = Pt(12)


def add_table_slide(prs, page):
    """添加表格页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题
    txBox = slide.shapes.add_textbox(Cm(2), Cm(1.5), Cm(17), Cm(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = page["title"]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE

    # 表格
    headers = page["headers"]
    rows = page["rows"]
    n_rows = len(rows) + 1
    n_cols = len(headers)

    table_left = Cm(2)
    table_top = Cm(5)
    table_width = Cm(17)
    table_height = Cm(1.5 * n_rows)

    table_shape = slide.shapes.add_table(n_rows, n_cols, table_left, table_top, table_width, table_height)
    table = table_shape.table

    # 表头
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(14)
            paragraph.font.bold = True
            paragraph.font.color.rgb = COLOR_WHITE
            paragraph.alignment = PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_BG

    # 数据行
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.color.rgb = COLOR_BODY
                paragraph.alignment = PP_ALIGN.CENTER
            # 交替行背景
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)


def create_pptx():
    """生成 PPTX 文件"""
    prs = Presentation()
    prs.slide_width = A4_WIDTH
    prs.slide_height = A4_HEIGHT

    for page in CONTENT["pages"]:
        page_type = page.get("type", "text")
        if page_type == "title":
            add_title_slide(prs, page)
        elif page_type == "text":
            add_text_slide(prs, page)
        elif page_type == "table":
            add_table_slide(prs, page)

    prs.save(PPTX_PATH)
    print(f"[OK] PPTX 已生成: {PPTX_PATH}")
    return PPTX_PATH


def pptx_to_pdf():
    """用 PowerPoint COM 将 PPTX 转为 PDF"""
    import comtypes.client

    # COM 需要绝对路径
    pptx_abs = os.path.abspath(PPTX_PATH)
    pdf_abs = os.path.abspath(PDF_PATH)

    print("[..] 启动 PowerPoint COM...")
    powerpoint = comtypes.client.CreateObject("PowerPoint.Application")
    # Visible: 1 = 可见, 0 = 隐藏（调试时建议设 1）
    powerpoint.Visible = 1

    try:
        deck = powerpoint.Presentations.Open(pptx_abs, WithWindow=False)
        # 32 = ppSaveAsPDF
        deck.SaveAs(pdf_abs, 32)
        deck.Close()
        print(f"[OK] PDF 已生成: {pdf_abs}")
    finally:
        powerpoint.Quit()


def main():
    print("=" * 50)
    print("方案一：PPT → PDF")
    print("=" * 50)

    # Step 1: 生成 PPTX
    create_pptx()

    # Step 2: 转为 PDF
    pptx_to_pdf()

    print("\n完成！")
    print(f"  PPTX: {PPTX_PATH}")
    print(f"  PDF:  {PDF_PATH}")


if __name__ == "__main__":
    main()
